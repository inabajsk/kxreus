# -*- coding: utf-8 -*-
"""Small python-pptx layer shared by build_deck.py (16:9, 13.333 x 7.5 in)."""
import os

from lxml import etree
from pptx import Presentation
from pptx.chart.data import XyChartData, CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x17, 0x22, 0x33)
ORANGE = RGBColor(0xEE, 0x7B, 0x1E)
TEAL = RGBColor(0x1F, 0x7A, 0x8C)
TILE = RGBColor(237, 242, 247)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x5B, 0x67, 0x78)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
CODEBG = RGBColor(0x1E, 0x27, 0x36)
CODEFG = RGBColor(0xE6, 0xED, 0xF3)
CODECM = RGBColor(0x8F, 0xB3, 0x8A)
SOFTOR = RGBColor(0xFD, 0xEB, 0xDB)
SOFTTL = RGBColor(0xDD, 0xEE, 0xF1)
PALE = RGBColor(0xC5, 0xD0, 0xDE)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)
# categorical series colours (validated for distinguishability on white)
SERIES = ["1F77B4", "EE7B1E", "2CA02C", "D62728", "9467BD", "8C564B",
          "E377C2", "7F7F7F", "BCBD22", "17BECF", "393B79", "AD494A"]

FONT = "Yu Gothic"
MONO = "Courier New"


class Deck:
  def __init__(self):
    self.prs = Presentation()
    self.prs.slide_width = Inches(13.333)
    self.prs.slide_height = Inches(7.5)
    self.blank = self.prs.slide_layouts[6]
    self.page = 0

  def slide(self, dark=False):
    self.page += 1
    s = self.prs.slides.add_slide(self.blank)
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = INK if dark else WHITE
    if self.page > 1:
      tb(s, 12.3, 7.05, 0.8, 0.3, [[{"t": str(self.page), "size": 10,
                                      "color": RGBColor(0x9A, 0xA8, 0xBA) if dark else MUTED}]],
         align=PP_ALIGN.RIGHT, margin=0)
    return s

  def save(self, path):
    self.prs.save(path)


def set_font(run, size, color=TEXT, bold=False, font=FONT, italic=False):
  run.font.size = Pt(size)
  run.font.bold = bold
  run.font.italic = italic
  run.font.color.rgb = color
  run.font.name = font
  rPr = run._r.get_or_add_rPr()
  for tag in ("a:ea", "a:cs"):
    el = rPr.find(qn(tag))
    if el is None:
      el = etree.SubElement(rPr, qn(tag))
    el.set("typeface", FONT if font == MONO else font)


def tb(slide, x, y, w, h, paras, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, font=FONT, margin=0.05, spacing=None, line=None):
  box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
  tf = box.text_frame
  tf.word_wrap = True
  tf.vertical_anchor = anchor
  for side in ("left", "right", "top", "bottom"):
    setattr(tf, "margin_" + side, Inches(margin))
  if isinstance(paras, str):
    paras = [paras]
  for i, p in enumerate(paras):
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    para.alignment = align
    if spacing is not None:
      para.space_after = Pt(spacing)
    if line is not None:
      para.line_spacing = line
    runs = p if isinstance(p, list) else [{"t": p}]
    for r in runs:
      run = para.add_run()
      run.text = r["t"]
      set_font(run, r.get("size", size), r.get("color", color), r.get("bold", bold),
               r.get("font", font), r.get("italic", False))
  return box


def bullets(slide, x, y, w, h, items, size=15, color=TEXT, gap=6, bullet_color=ORANGE):
  box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
  tf = box.text_frame
  tf.word_wrap = True
  for side in ("left", "right", "top", "bottom"):
    setattr(tf, "margin_" + side, Inches(0.03))
  for i, it in enumerate(items):
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    level = 0
    if isinstance(it, tuple):
      level, it = it
    runs = it if isinstance(it, list) else [{"t": it}]
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.22 + 0.25 * level))))
    pPr.set("indent", str(int(-Inches(0.2))))
    buClr = etree.SubElement(pPr, qn("a:buClr"))
    clr = etree.SubElement(buClr, qn("a:srgbClr"))
    clr.set("val", "%02X%02X%02X" % tuple(bullet_color if level == 0 else MUTED))
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", "■" if level == 0 else "–")
    para.space_after = Pt(gap)
    for r in runs:
      run = para.add_run()
      run.text = r["t"]
      set_font(run, r.get("size", size - (1 if level else 0)), r.get("color", color),
               r.get("bold", False), r.get("font", FONT))
  return box


def rect(slide, x, y, w, h, fill=TILE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.08):
  s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
  s.fill.solid()
  s.fill.fore_color.rgb = fill
  if line is None:
    s.line.fill.background()
  else:
    s.line.color.rgb = line
    s.line.width = Pt(1.25)
  if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
    s.adjustments[0] = radius
  st = s._element.find(qn("p:style"))  # theme style draws a shadow
  if st is not None:
    s._element.remove(st)
  s.text_frame.text = ""
  return s


def shape_text(s, text, size=14, color=TEXT, bold=False, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, font=FONT, margin=0.06):
  tf = s.text_frame
  tf.word_wrap = True
  tf.vertical_anchor = anchor
  for side in ("left", "right", "top", "bottom"):
    setattr(tf, "margin_" + side, Inches(margin))
  lines = text if isinstance(text, list) else [text]
  for i, ln in enumerate(lines):
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    para.alignment = align
    runs = ln if isinstance(ln, list) else [{"t": ln}]
    for r in runs:
      run = para.add_run()
      run.text = r["t"]
      set_font(run, r.get("size", size), r.get("color", color), r.get("bold", bold), r.get("font", font))


def box(slide, x, y, w, h, t1, t2=None, fill=TILE, c1=INK, c2=MUTED, s1=13, s2=10.5, line=None):
  r = rect(slide, x, y, w, h, fill, line=line)
  lines = [[{"t": t1, "bold": True, "color": c1, "size": s1}]]
  if t2:
    for t in (t2 if isinstance(t2, list) else [t2]):
      lines.append([{"t": t, "color": c2, "size": s2}])
  shape_text(r, lines)
  return r


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0, head=True):
  c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
  c.line.color.rgb = color
  c.line.width = Pt(width)
  if head:
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
  return c


def code(slide, x, y, w, h, lines, size=12):
  rect(slide, x, y, w, h, CODEBG, radius=0.04)
  paras = []
  for ln in lines:
    if "#" in ln and not ln.startswith("$"):
      i = ln.index("#")
      runs = ([{"t": ln[:i], "color": CODEFG}] if ln[:i] else []) + [{"t": ln[i:], "color": CODECM}]
      paras.append(runs)
    elif ln.startswith("$"):
      paras.append([{"t": ln, "color": RGBColor(0xF7, 0xC0, 0x8A)}])
    else:
      paras.append([{"t": ln if ln else " ", "color": CODEFG}])
  tb(slide, x + 0.15, y + 0.12, w - 0.3, h - 0.24, paras, size=size, font=MONO, margin=0.02, line=1.05)


def title(slide, t, kicker=None, dark=False, size=28):
  if kicker:
    tb(slide, 0.6, 0.32, 12, 0.35, [[{"t": kicker, "bold": True, "color": ORANGE, "size": 12.5}]], margin=0)
  tb(slide, 0.6, 0.6, 12.2, 0.8, [[{"t": t, "bold": True, "color": WHITE if dark else INK, "size": size}]], margin=0)


def notes(slide, text):
  slide.notes_slide.notes_text_frame.text = text


def image(slide, path, x, y, w, h=None, tile=False):
  if tile:
    rect(slide, x, y, w, h or w, TILE, radius=0.06)
  if not os.path.exists(path):
    return None
  if h is None:
    return slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
  return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))


def movie(slide, path, poster, x, y, w, h):
  if not os.path.exists(path):
    r = rect(slide, x, y, w, h, TILE)
    shape_text(r, [[{"t": "（動画：計測中）", "color": MUTED, "size": 14}]])
    return None
  return slide.shapes.add_movie(path, Inches(x), Inches(y), Inches(w), Inches(h),
                                poster_frame_image=poster if poster and os.path.exists(poster) else None,
                                mime_type="video/mp4")


def table(slide, x, y, w, rows, colw, size=11, rowh=0.36, head_fill=INK, zebra=True,
          mono_col0=False, align_right_from=None):
  nr, nc = len(rows), len(rows[0])
  gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
  t = gt.table
  tblPr = t._tbl.tblPr
  tblPr.set("bandRow", "0")
  tblPr.set("firstRow", "0")
  for j, cw in enumerate(colw):
    t.columns[j].width = Inches(cw)
  for i in range(nr):
    t.rows[i].height = Inches(rowh)
    for j in range(nc):
      cell = t.cell(i, j)
      cell.margin_left = Inches(0.07)
      cell.margin_right = Inches(0.07)
      cell.margin_top = Inches(0.02)
      cell.margin_bottom = Inches(0.02)
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.fill.solid()
      cell.fill.fore_color.rgb = head_fill if i == 0 else (TILE if (zebra and i % 2 == 0) else WHITE)
      tf = cell.text_frame
      tf.word_wrap = True
      para = tf.paragraphs[0]
      if align_right_from is not None and j >= align_right_from and i > 0:
        para.alignment = PP_ALIGN.RIGHT
      val = rows[i][j]
      runs = val if isinstance(val, list) else [{"t": str(val)}]
      for r in runs:
        run = para.add_run()
        run.text = r["t"]
        if i == 0:
          set_font(run, r.get("size", size), WHITE, True)
        else:
          f = MONO if (j == 0 and mono_col0) else r.get("font", FONT)
          set_font(run, r.get("size", size), r.get("color", TEXT), r.get("bold", j == 0 and mono_col0), f)
  return gt


def xy_chart(slide, x, y, w, h, series, title_text=None, x_title=None, y_title=None,
             colors=None, legend=True, y_min=None, y_max=None, font_size=10, x_max=None):
  """series: list of (name, [(x, y), ...])."""
  cd = XyChartData()
  for name, pts in series:
    s = cd.add_series(name)
    for a, b in pts:
      s.add_data_point(a, b)
  gf = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS, Inches(x), Inches(y),
                              Inches(w), Inches(h), cd)
  ch = gf.chart
  ch.font.size = Pt(font_size)
  ch.font.name = FONT
  ch.has_title = bool(title_text)
  if title_text:
    ch.chart_title.text_frame.text = title_text
    for p in ch.chart_title.text_frame.paragraphs:
      for r in p.runs:
        set_font(r, font_size + 2, TEXT, True)
  ch.has_legend = legend
  if legend:
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(font_size - 1)
  for i, s in enumerate(ch.series):
    col = (colors or SERIES)[i % len(colors or SERIES)]
    s.format.line.color.rgb = RGBColor.from_string(col)
    s.format.line.width = Pt(1.75)
    s.smooth = False
  va = ch.value_axis
  va.has_major_gridlines = True
  va.major_gridlines.format.line.color.rgb = LINE
  va.format.line.color.rgb = LINE
  va.tick_labels.font.size = Pt(font_size)
  if y_min is not None:
    va.minimum_scale = y_min
  if y_max is not None:
    va.maximum_scale = y_max
  ca = ch.category_axis
  ca.has_major_gridlines = False
  ca.format.line.color.rgb = LINE
  ca.tick_labels.font.size = Pt(font_size)
  ca.minimum_scale = 0
  if x_max is not None:
    ca.maximum_scale = x_max
  for ax, t in ((ca, x_title), (va, y_title)):
    if t:
      ax.has_title = True
      ax.axis_title.text_frame.text = t
      for p in ax.axis_title.text_frame.paragraphs:
        for r in p.runs:
          set_font(r, font_size, MUTED)
  return ch


def bar_chart(slide, x, y, w, h, categories, series, colors=None, title_text=None, y_max=None,
              font_size=10, number_format='0%', legend=True, horizontal=False, labels=True):
  cd = CategoryChartData()
  cd.categories = categories
  for name, vals in series:
    cd.add_series(name, vals)
  kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
  gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
  ch = gf.chart
  ch.font.size = Pt(font_size)
  ch.font.name = FONT
  ch.has_title = bool(title_text)
  if title_text:
    ch.chart_title.text_frame.text = title_text
    for p in ch.chart_title.text_frame.paragraphs:
      for r in p.runs:
        set_font(r, font_size + 2, TEXT, True)
  ch.has_legend = legend and len(series) > 1
  if ch.has_legend:
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
  plot = ch.plots[0]
  plot.gap_width = 60
  plot.overlap = -10
  if labels:
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(font_size - 1)
  for i, s in enumerate(plot.series):
    s.format.fill.solid()
    s.format.fill.fore_color.rgb = RGBColor.from_string((colors or SERIES)[i % len(colors or SERIES)])
  va = ch.value_axis
  va.has_major_gridlines = True
  va.major_gridlines.format.line.color.rgb = LINE
  va.format.line.color.rgb = LINE
  va.tick_labels.number_format = number_format
  va.tick_labels.number_format_is_linked = False
  va.minimum_scale = 0
  if y_max is not None:
    va.maximum_scale = y_max
  ch.category_axis.format.line.color.rgb = LINE
  ch.category_axis.tick_labels.font.size = Pt(font_size)
  return ch
